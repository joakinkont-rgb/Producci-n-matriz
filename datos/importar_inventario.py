import openpyxl, json, re, unicodedata

SRC = "/root/.claude/uploads/2f7c318a-942b-54c9-bc63-bd9a2c8772a1/c8e655cf-INV_MATRIZ.xlsx"
wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["Hoja1"]

# (titulo_en_archivo, fila_inicio, fila_fin, clave_almacen, nombre_almacen, zona)
SECCIONES = [
    ("BODEGA",        4,  71, "bodega",      "Bodega de secos",             "secos"),
    ("QUESOS",       75,  85, "refrigerado", "Refrigerado · Lácteos y varios","refrigerado"),
    ("LIMPIEZA",     89,  95, "limpieza",    "Limpieza",                    "noalimento"),
    ("VERDURA",     100, 134, "verdura",     "Cuarto frío · Verdura",       "verdura"),
    ("DESHECHABLES",138, 157, "desechables", "Desechables",                 "noalimento"),
    ("CAFETERIAS",  161, 182, "cafeterias",  "Congelador · Al vacío",       "congelador"),
]

def slug(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    s = re.sub(r"[^A-Za-z0-9]+", "-", s).strip("-").lower()
    return s or "x"

articulos, unidades, resumen = [], {}, []
for titulo, r0, r1, almacen, nom_alm, zona in SECCIONES:
    n = 0
    for r in range(r0, r1 + 1):
        nombre = ws.cell(row=r, column=2).value
        if nombre is None or not str(nombre).strip():
            continue
        nombre = str(nombre).strip()
        mn = ws.cell(row=r, column=3).value
        u  = ws.cell(row=r, column=6).value
        nota = ws.cell(row=r, column=7).value
        u = (str(u).strip() if u is not None and str(u).strip() else None)
        if u:
            u = {"SACOS": "SACO", "CUBET": "CUB", "PIEZA": "PZ"}.get(u.upper(), u.upper())
        if u:
            unidades[u] = unidades.get(u, 0) + 1
        articulos.append({
            "id": f"{almacen}-{slug(nombre)}",
            "codigo": None,
            "nombre": nombre,
            "tipo": "terminado" if almacen == "cafeterias" else "insumo",
            "categoria": titulo.capitalize(),
            "unidad": u if u else "PZ",
            "almacen": almacen,
            "stockMin": float(mn) if isinstance(mn, (int, float)) else None,
            "vidaUtilDias": None, "costoRef": None, "activo": True,
            "nota": str(nota).strip() if nota else None,
            "origenFila": r,
        })
        n += 1
    resumen.append((titulo, nom_alm, n))

almacenes = [{"id": a, "nombre": nom, "zona": z, "temp": t, "orden": i + 1}
             for i, (_, _, _, a, nom, z, t) in enumerate(
                 [(s[0], s[1], s[2], s[3], s[4], s[5], t) for s, t in zip(SECCIONES, [
                     "Ambiente", "2 a 6 °C", "Ambiente", "2 a 6 °C", "Ambiente", "−18 °C"])])]
# almacén de cárnicos: mencionado por el usuario, ausente del archivo
almacenes.append({"id": "carnico", "nombre": "Cuarto frío · Cárnicos", "zona": "carnico",
                  "temp": "0 a 4 °C", "orden": 7})

json.dump({"almacenes": almacenes, "articulos": articulos},
          open("/home/user/Producci-n-matriz/datos/inventario-matriz.json", "w"),
          ensure_ascii=False, indent=1)

print("SECCIONES IMPORTADAS")
for t, nom, n in resumen:
    print(f"  {t:<14} -> {nom:<32} {n:>3} artículos")
print(f"\nTOTAL artículos: {len(articulos)}   almacenes: {len(almacenes)}")
print("\nUNIDADES encontradas (unidad: veces):")
for u, c in sorted(unidades.items(), key=lambda x: -x[1]):
    print(f"  {u:<10} {c}")
sin = [a["nombre"] for a in articulos if a["stockMin"] is None]
print(f"\nSin mínimo definido ({len(sin)}): {', '.join(sin)}")
